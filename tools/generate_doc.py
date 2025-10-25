import os
import io
import sys
from datetime import datetime

try:
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
except Exception as e:
    print('Missing python-docx. Please install it and rerun: pip install python-docx')
    raise

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
EXCLUDE_DIRS = {'.git', '__pycache__', '.venv', 'venv', 'node_modules', '.pytest_cache', '.idea'}

# Simple purpose hints for common files
PURPOSE_HINTS = {
    'manage.py': 'Django management entrypoint (runserver, migrations, etc.)',
    'requirements.txt': 'Pinned Python dependencies for the project',
    'db.sqlite3': 'SQLite database file (development)',
    'artshare/settings.py': 'Django settings and configuration (including storage)',
    'artshare/urls.py': 'Project URL routing',
    'artshare/wsgi.py': 'WSGI entrypoint for WSGI servers',
    'artshare/asgi.py': 'ASGI entrypoint for async servers',
    'core/models.py': 'Django models: Artist and Artwork',
    'core/views.py': 'API views (registration, login, artwork endpoints)',
    'core/serializers.py': 'DRF serializers for Artist and Artwork',
    'core/urls.py': 'App-level URL routing for core APIs',
}

# Files to annotate more deeply
ANNOTATE_FILES = {
    os.path.join('artshare', 'settings.py'),
    os.path.join('core', 'models.py'),
    os.path.join('core', 'views.py'),
    os.path.join('core', 'serializers.py'),
    'requirements.txt',
}


def should_exclude(path):
    parts = set(path.split(os.sep))
    return bool(parts & EXCLUDE_DIRS)


def collect_files(root):
    files = []
    for dirpath, dirnames, filenames in os.walk(root):
        # filter dirnames in-place
        dirnames[:] = [d for d in dirnames if d not in EXCLUDE_DIRS]
        for fn in filenames:
            # skip compiled and binary files
            if fn.endswith(('.pyc', '.pyo', '.so', '.dll', '.exe')):
                continue
            full = os.path.join(dirpath, fn)
            rel = os.path.relpath(full, root)
            files.append((rel, full))
    files.sort()
    return files


def short_purpose(relpath):
    return PURPOSE_HINTS.get(relpath, 'Source file or asset. See contents for details.')


def extract_key_symbols_py(content):
    symbols = []
    for line in content.splitlines():
        line = line.strip()
        if line.startswith('class '):
            symbols.append(line.split(':')[0])
        if line.startswith('def '):
            symbols.append(line.split('(')[0])
    return symbols


def annotate_settings(content):
    notes = []
    lines = content.splitlines()
    for i, l in enumerate(lines):
        if 'SECRET_KEY' in l:
            notes.append(('SECRET_KEY', 'Secret used for cryptographic signing; must be kept private in production.'))
        if l.strip().startswith('DEBUG'):
            notes.append(('DEBUG', 'Toggle debug mode. Should be False in production.'))
        if 'ALLOWED_HOSTS' in l:
            notes.append(('ALLOWED_HOSTS', 'Hosts/domains the Django site can serve.'))
        if 'INSTALLED_APPS' in l:
            notes.append(('INSTALLED_APPS', 'List of enabled Django apps; cloudinary and cloudinary_storage control remote media.'))
        if 'CLOUDINARY' in l:
            notes.append(('CLOUDINARY', 'Cloudinary env vars and configuration for remote storage.'))
        if 'DEFAULT_FILE_STORAGE' in l:
            notes.append(('DEFAULT_FILE_STORAGE', 'Storage backend used for user-uploaded files (Cloudinary or FileSystem).'))
        if 'MEDIA_ROOT' in l or 'MEDIA_URL' in l:
            notes.append(('MEDIA', 'Local media settings; used when DEFAULT_FILE_STORAGE is FileSystemStorage.'))
        if 'DATABASES' in l:
            notes.append(('DATABASES', 'Database configuration, parsed from DATABASE_URL (Render).'))
    # Deduplicate
    seen = set()
    out = []
    for k, v in notes:
        if k not in seen:
            out.append((k, v))
            seen.add(k)
    return out


def annotate_models(content):
    notes = []
    if 'class Artist' in content:
        notes.append(('Artist', 'Custom user model extending AbstractUser; contains bio and profile_picture ImageField.'))
    if 'class Artwork' in content:
        notes.append(('Artwork', 'Artwork model with ForeignKey to Artist and an ImageField for the artwork image.'))
    if "upload_to='" in content or 'upload_to="' in content:
        notes.append(('upload_to', 'Determines the destination path/prefix for uploaded files. When using Cloudinary storage, UPLOAD_OPTIONS.folder controls folder on Cloudinary.'))
    return notes


def annotate_views(content):
    notes = []
    if 'RegisterArtistView' in content:
        notes.append(('RegisterArtistView', 'Handles registration; accepts multipart for profile pictures.'))
    if 'ArtworkListCreateView' in content:
        notes.append(('ArtworkListCreateView', 'List and create endpoint for artworks; ensures authenticated uploads.'))
    if 'parser_classes' in content:
        notes.append(('parsers', 'MultiPartParser/FormParser enabled to accept file uploads via multipart/form-data.'))
    return notes


def annotate_serializers(content):
    notes = []
    if 'ArtistSerializer' in content:
        notes.append(('ArtistSerializer', 'Handles creation and update of Artist; manages password hashing and profile_picture.'))
    if 'ArtworkSerializer' in content:
        notes.append(('ArtworkSerializer', 'Serializes Artwork; create() uses request.user and saves image through storage backend.'))
    return notes


def add_code_block(doc, text):
    # Add code block with monospace font
    # sanitize control characters that python-docx / lxml can't accept
    def sanitize_line(s):
        return ''.join(ch if (ord(ch) >= 32 or ord(ch) in (9, 10, 13)) else ' ' for ch in s)

    # Add content line-by-line to avoid very large single runs and to reduce
    # the chance of problematic characters breaking lxml.
    for line in text.splitlines():
        safe = sanitize_line(line)
        p = doc.add_paragraph()
        run = p.add_run(safe)
        run.font.name = 'Courier New'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')
        run.font.size = Pt(8)


def build_doc(output_path):
    doc = Document()
    doc.core_properties.title = 'Art-Share Django Project Documentation'
    doc.core_properties.created = datetime.utcnow()

    doc.add_heading('Art-Share Django Project', level=1)
    doc.add_paragraph(f'Generated: {datetime.utcnow().isoformat()} UTC')
    doc.add_paragraph('Summary: This document describes the repository structure, important files, and focused annotations for critical Django modules (settings, models, views, serializers, requirements).')

    files = collect_files(ROOT)

    for rel, full in files:
        doc.add_heading(rel, level=2)
        purpose = short_purpose(rel)
        doc.add_paragraph(f'Purpose: {purpose}')
        try:
            with open(full, 'r', encoding='utf-8') as f:
                content = f.read()
        except Exception as e:
            doc.add_paragraph(f'Could not read file: {e}')
            continue

        # Key symbols for python files
        if rel.endswith('.py'):
            symbols = extract_key_symbols_py(content)
            if symbols:
                doc.add_paragraph('Key symbols: ' + ', '.join(symbols))

        # Focused annotations for selected files
        if rel in ANNOTATE_FILES or rel.replace('\\','/') in ANNOTATE_FILES:
            doc.add_paragraph('Annotations:')
            if rel.endswith('settings.py') or rel == os.path.join('artshare', 'settings.py'):
                for k, v in annotate_settings(content):
                    doc.add_paragraph(f'- {k}: {v}')
            elif rel.endswith('models.py'):
                for k, v in annotate_models(content):
                    doc.add_paragraph(f'- {k}: {v}')
            elif rel.endswith('views.py'):
                for k, v in annotate_views(content):
                    doc.add_paragraph(f'- {k}: {v}')
            elif rel.endswith('serializers.py'):
                for k, v in annotate_serializers(content):
                    doc.add_paragraph(f'- {k}: {v}')
            elif rel == 'requirements.txt':
                doc.add_paragraph('- Requirements: List of Python packages installed for the project')

        doc.add_paragraph('Contents:')
        # Add the file content as a code block (monospace)
        try:
            add_code_block(doc, content)
        except Exception as e:
            # If adding the raw content fails (control characters, etc.),
            # include a sanitized placeholder and skip the full dump.
            doc.add_paragraph('Could not include full file contents due to encoding/control characters. Showing sanitized excerpt:')
            safe = ''.join(ch if (ord(ch) >= 32 or ord(ch) in (9,10,13)) else ' ' for ch in content)
            excerpt = safe[:2000]
            add_code_block(doc, excerpt)
        # Add a page break between files for readability
        doc.add_page_break()

    doc.save(output_path)


if __name__ == '__main__':
    out = os.path.join(ROOT, 'artshare-project-documentation.docx')
    print('Generating', out)
    build_doc(out)
    print('Done')
